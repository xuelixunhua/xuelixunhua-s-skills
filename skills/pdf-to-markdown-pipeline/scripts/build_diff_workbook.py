#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build an Excel workbook for clause diff + parameter diff + linkage view.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SUBHEADER_FILL = PatternFill("solid", fgColor="EAF4E2")
NOTE_FILL = PatternFill("solid", fgColor="FFF7D6")
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

CATEGORY_LABELS = {
    "price_limit": "价格边界",
    "time_point": "时间点",
    "duration": "时长窗口",
    "count_limit": "次数限制",
    "ratio_threshold": "比例阈值",
    "power_capacity": "功率/容量",
    "penalty_factor": "罚因子",
    "voltage_level": "电压等级",
    "generic_numeric_rule": "一般数值规则",
}


def resolve_path(raw_path: str) -> Path:
    path = Path(raw_path).expanduser()
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()
    return path


def refresh_frontend_dataset() -> None:
    root = next(
        (
            candidate
            for candidate in [Path(__file__).resolve(), *Path(__file__).resolve().parents]
            if (candidate / "00_总控").exists() and (candidate / "1_Inputs").exists()
        ),
        Path(__file__).resolve().parents[3],
    )
    builder = root / "0_脚本" / "个性化能力" / "前端装配" / "build_frontend_dataset.py"
    if not builder.exists():
        print("Frontend dataset builder not found, skipped refresh.")
        return

    result = subprocess.run(
        [sys.executable, str(builder)],
        cwd=root,
        check=False,
    )
    if result.returncode != 0:
        raise SystemExit("Frontend dataset refresh failed after workbook generation.")


def write_header(sheet, row_idx: int, values: List[str], fill=HEADER_FILL) -> None:
    for col_idx, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_idx, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = WRAP_ALIGNMENT


def autosize(sheet, widths: Dict[int, int]) -> None:
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width


def finalize_sheet(sheet, freeze_panes: Optional[str] = "A2") -> None:
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    if freeze_panes:
        sheet.freeze_panes = freeze_panes


def normalize_text(text: str) -> str:
    text = text or ""
    text = re.sub(r"\s+", "", text)
    return text


def yes_no(value: bool) -> str:
    return "是" if value else "否"


def category_label(raw_category: str) -> str:
    return CATEGORY_LABELS.get(raw_category or "", raw_category or "")


def safe_clause_text(row: Dict, field: str) -> str:
    if row.get("formula_sensitive"):
        markdown_field = {
            "old_text": "old_markdown",
            "new_text": "new_markdown",
            "text": "markdown",
        }.get(field)
        if markdown_field and row.get(markdown_field):
            return row.get(markdown_field, "")
        return "该条款涉及复杂公式，但当前未提取到可复用的 LaTeX/Markdown，请结合原图留痕与原文复核。"
    return row.get(field, "")


def display_rule_title(document: Dict) -> str:
    title = (document or {}).get("title") or ""
    if title:
        return title
    source_path = (document or {}).get("source_path") or ""
    if source_path:
        return Path(source_path).stem
    return "规则"


def infer_subject(*texts: str) -> str:
    merged = "".join(texts)
    if any(token in merged for token in ("容量电费考核", "发电能力抽查")):
        return "容量电费考核"
    if any(token in merged for token in ("独立新型储能", "独立储能", "储能电站", "储能设施", "荷电状态", "充放电")):
        return "独立储能"
    if "新能源" in merged:
        return "新能源"
    if any(token in merged for token in ("火电", "机组", "启动费用")):
        return "火电机组"
    if any(token in merged for token in ("断面", "安全约束", "控制裕度")):
        return "电网安全约束"
    if any(token in merged for token in ("运行备用", "备用容量")):
        return "运行备用"
    if any(token in merged for token in ("价格", "报价", "出清")):
        return "价格机制"
    if "附录3" in merged:
        return "机组运行参数"
    if "附录4" in merged:
        return "默认申报参数"
    if "附录5" in merged:
        return "市场核定参数"
    return "待人工判定"


def infer_change_detail(change_type: str, old_value: str, new_value: str) -> str:
    if change_type == "新增参数":
        return "规则新增"
    if change_type == "删除参数":
        return "规则移除"
    if any(token in f"{old_value}{new_value}" for token in ("-", "至", "~", "～")):
        return "边界调整"
    return "数值调整"


def infer_business_impact(
    raw_category: str,
    subject: str,
    source_ref: str,
    statement_text: str,
) -> str:
    merged = f"{subject}{source_ref}{statement_text}"
    if subject == "独立储能":
        return "影响独立储能参与方式、循环约束或充放电边界"
    if subject == "容量电费考核":
        return "影响容量电费考核节奏、申报窗口或执行要求"
    if raw_category == "price_limit":
        return "影响申报边界、价格空间或报价约束"
    if raw_category == "ratio_threshold" or subject == "电网安全约束":
        return "影响安全校核、控制裕度或触发阈值"
    if raw_category == "count_limit":
        return "影响启停次数或循环次数限制"
    if raw_category in {"duration", "time_point"}:
        return "影响执行时点、持续时长或调度节奏"
    if raw_category in {"power_capacity", "voltage_level"}:
        return "影响出力、容量、电压或边界条件"
    if "附录5" in merged:
        return "影响市场核定参数与市场边界条件"
    if "附录3" in merged:
        return "影响机组运行约束与可行域边界"
    return "建议结合条款语境进一步判断"


def is_key_change(
    raw_category: str,
    subject: str,
    source_ref: str,
    statement_text: str,
) -> bool:
    merged = f"{subject}{source_ref}{statement_text}"
    if subject in {"独立储能", "容量电费考核", "电网安全约束", "运行备用", "价格机制"}:
        return True
    if raw_category in {"price_limit", "ratio_threshold", "count_limit", "penalty_factor"}:
        return True
    if any(token in merged for token in ("申报价格", "出清价格", "控制裕度", "断面限值", "启动费用")):
        return True
    return False


def needs_manual_review(
    source_type: str,
    raw_category: str,
    statement_text: str,
) -> bool:
    if source_type != "clause":
        return False
    if raw_category in {"time_point", "duration", "generic_numeric_rule"}:
        return True
    if any(token in statement_text for token in ("T-", "T+", "D日", "D-1日", "次日", "后一日")):
        return True
    return False


def build_clause_index(structured: Dict) -> Dict[str, Dict]:
    result: Dict[str, Dict] = {}
    for clause in structured.get("clauses", []):
        key = f"{clause['marker']}|{clause.get('label') or ''}"
        result[key] = clause
    return result


def build_parameter_row(
    change_type: str,
    parameter_name: str,
    context_label: str,
    old_value: str,
    new_value: str,
    unit: str,
    raw_category: str,
    source_type: str,
    old_source_ref: str,
    new_source_ref: str,
    statement_text: str,
) -> Dict:
    subject = infer_subject(parameter_name, context_label, old_source_ref, new_source_ref, statement_text)
    key_change = is_key_change(raw_category, subject, new_source_ref or old_source_ref, statement_text)
    manual_review = needs_manual_review(source_type, raw_category, statement_text)
    return {
        "change_type": change_type,
        "change_detail": infer_change_detail(change_type, old_value, new_value),
        "raw_category": raw_category,
        "category_label": category_label(raw_category),
        "subject": subject,
        "is_key_change": yes_no(key_change),
        "should_report": yes_no(key_change),
        "needs_manual_review": yes_no(manual_review),
        "parameter_name": parameter_name,
        "context_label": context_label,
        "old_value": old_value,
        "new_value": new_value,
        "unit": unit,
        "old_source_ref": old_source_ref,
        "new_source_ref": new_source_ref,
        "business_impact": infer_business_impact(raw_category, subject, new_source_ref or old_source_ref, statement_text),
        "statement_text": statement_text,
        "source_type": source_type,
    }


def collect_parameter_rows(parameter_diff: Dict) -> List[Dict]:
    rows: List[Dict] = []

    for item in parameter_diff.get("changed_items", []):
        old_item = item["old_items"][0]
        new_item = item["new_items"][0]
        rows.append(
            build_parameter_row(
                change_type="数值变化",
                parameter_name=item["parameter_name"],
                context_label=item["context_label"],
                old_value=" / ".join(x["value_text"] for x in item["old_items"]),
                new_value=" / ".join(x["value_text"] for x in item["new_items"]),
                unit=new_item.get("unit") or old_item.get("unit") or "",
                raw_category=new_item.get("parameter_category") or old_item.get("parameter_category") or "",
                source_type=new_item.get("source_type") or old_item.get("source_type") or "",
                old_source_ref=old_item.get("source_ref", ""),
                new_source_ref=new_item.get("source_ref", ""),
                statement_text=new_item.get("display_statement_text")
                or old_item.get("display_statement_text")
                or new_item.get("statement_text")
                or old_item.get("statement_text")
                or "",
            )
        )

    for item in parameter_diff.get("added_items", []):
        rows.append(
            build_parameter_row(
                change_type="新增参数",
                parameter_name=item["parameter_name"],
                context_label=item["context_label"],
                old_value="",
                new_value=item["value_text"],
                unit=item.get("unit") or "",
                raw_category=item.get("parameter_category") or "",
                source_type=item.get("source_type") or "",
                old_source_ref="",
                new_source_ref=item.get("source_ref", ""),
                statement_text=item.get("display_statement_text") or item.get("statement_text") or "",
            )
        )

    for item in parameter_diff.get("removed_items", []):
        rows.append(
            build_parameter_row(
                change_type="删除参数",
                parameter_name=item["parameter_name"],
                context_label=item["context_label"],
                old_value=item["value_text"],
                new_value="",
                unit=item.get("unit") or "",
                raw_category=item.get("parameter_category") or "",
                source_type=item.get("source_type") or "",
                old_source_ref=item.get("source_ref", ""),
                new_source_ref="",
                statement_text=item.get("display_statement_text") or item.get("statement_text") or "",
            )
        )

    return rows


def collect_linkage_rows(
    old_structured: Dict,
    new_structured: Dict,
    clause_diff: Dict,
    parameter_diff: Dict,
) -> List[Dict]:
    old_clause_index = build_clause_index(old_structured)
    new_clause_index = build_clause_index(new_structured)

    appendix_links: Dict[str, List[Dict]] = {}
    clause_links: Dict[str, List[Dict]] = {}

    def add_link(target: Dict[str, List[Dict]], key: str, payload: Dict) -> None:
        target.setdefault(key, []).append(payload)

    for item in parameter_diff.get("changed_items", []):
        old_item = item["old_items"][0]
        new_item = item["new_items"][0]
        payload = {
            "change_type": "数值变化",
            "parameter_name": item["parameter_name"],
            "context_label": item["context_label"],
            "old_value": " / ".join(x["value_text"] for x in item["old_items"]),
            "new_value": " / ".join(x["value_text"] for x in item["new_items"]),
            "source_type": new_item["source_type"],
            "source_ref": new_item["source_ref"],
            "parameter_category": new_item.get("parameter_category") or old_item.get("parameter_category") or "",
            "unit": new_item.get("unit") or old_item.get("unit") or "",
            "statement_text": new_item.get("display_statement_text")
            or old_item.get("display_statement_text")
            or new_item.get("statement_text")
            or old_item.get("statement_text")
            or "",
        }
        if new_item["source_type"] == "clause":
            marker = new_item["source_marker"]
            label = new_item["source_label"]
            add_link(clause_links, f"{marker}|{label}", payload)
        else:
            add_link(appendix_links, new_item["source_marker"], payload)

    for bucket_name, flag in (("added_items", "新增参数"), ("removed_items", "删除参数")):
        for item in parameter_diff.get(bucket_name, []):
            payload = {
                "change_type": flag,
                "parameter_name": item["parameter_name"],
                "context_label": item["context_label"],
                "old_value": "" if flag == "新增参数" else item["value_text"],
                "new_value": item["value_text"] if flag == "新增参数" else "",
                "source_type": item["source_type"],
                "source_ref": item["source_ref"],
                "parameter_category": item.get("parameter_category") or "",
                "unit": item.get("unit") or "",
                "statement_text": item.get("display_statement_text") or item.get("statement_text") or "",
            }
            if item["source_type"] == "clause":
                add_link(clause_links, f"{item['source_marker']}|{item['source_label']}", payload)
            else:
                add_link(appendix_links, item["source_marker"], payload)

    rows: List[Dict] = []
    for entry in clause_diff.get("revised_clauses", []) + clause_diff.get("added_clauses", []):
        marker = entry.get("new_marker") or entry.get("marker")
        label = entry.get("label") or ""
        clause_key = f"{marker}|{label}"
        clause_obj = new_clause_index.get(clause_key)
        if clause_obj is None and entry.get("old_marker"):
            clause_obj = old_clause_index.get(f"{entry['old_marker']}|{label}")

        linked_items: List[Dict] = []
        linked_items.extend(clause_links.get(clause_key, []))

        if clause_obj:
            for appendix_marker in clause_obj.get("referenced_appendices", []):
                linked_items.extend(appendix_links.get(appendix_marker.replace(" ", ""), []))
                linked_items.extend(appendix_links.get(appendix_marker, []))

        if not linked_items:
            continue

        seen = set()
        deduped = []
        for item in linked_items:
            dedupe_key = (
                item["change_type"],
                item["parameter_name"],
                item["context_label"],
                item.get("old_value", ""),
                item.get("new_value", ""),
                item["source_ref"],
            )
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            deduped.append(item)

        for item in deduped:
            subject = infer_subject(
                item["parameter_name"],
                item["context_label"],
                item["source_ref"],
                item.get("statement_text", ""),
                label,
            )
            rows.append(
                {
                    "chapter_title": entry.get("chapter_title", ""),
                    "clause_marker_old": entry.get("old_marker", ""),
                    "clause_marker_new": entry.get("new_marker", entry.get("marker", "")),
                    "clause_label": label,
                    "change_type": item["change_type"],
                    "change_detail": infer_change_detail(
                        item["change_type"],
                        item.get("old_value", ""),
                        item.get("new_value", ""),
                    ),
                    "category_label": category_label(item.get("parameter_category", "")),
                    "subject": subject,
                    "is_key_change": yes_no(
                        is_key_change(
                            item.get("parameter_category", ""),
                            subject,
                            item["source_ref"],
                            item.get("statement_text", ""),
                        )
                    ),
                    "needs_manual_review": yes_no(
                        needs_manual_review(
                            item["source_type"],
                            item.get("parameter_category", ""),
                            item.get("statement_text", ""),
                        )
                    ),
                    "parameter_name": item["parameter_name"],
                    "parameter_context": item["context_label"],
                    "parameter_old_value": item.get("old_value", ""),
                    "parameter_new_value": item.get("new_value", ""),
                    "parameter_source_ref": item["source_ref"],
                    "link_basis": "条款直接参数" if item["source_type"] == "clause" else "附录引用参数",
                    "business_impact": infer_business_impact(
                        item.get("parameter_category", ""),
                        subject,
                        item["source_ref"],
                        item.get("statement_text", ""),
                    ),
                }
            )
    return rows


def fill_summary(sheet, clause_diff: Dict, parameter_diff: Dict, linkage_rows: List[Dict]) -> None:
    sheet.title = "摘要"
    old_doc = clause_diff.get("old_document", {})
    new_doc = clause_diff.get("new_document", {})
    title = display_rule_title(new_doc) or display_rule_title(old_doc)
    old_version = old_doc.get("version_label", "")
    new_version = new_doc.get("version_label", "")
    version_part = f" {old_version}-{new_version}" if old_version and new_version else ""
    sheet["A1"] = f"{title}{version_part} 差异总览"
    sheet["A1"].font = Font(bold=True, size=14)

    write_header(sheet, 3, ["类别", "指标", "数值"])
    rows = [
        ("条款层", "旧版本条款数", clause_diff["summary"]["old_clause_count"]),
        ("条款层", "新版本条款数", clause_diff["summary"]["new_clause_count"]),
        ("条款层", "新增条款", clause_diff["summary"]["added_count"]),
        ("条款层", "删除条款", clause_diff["summary"]["removed_count"]),
        ("条款层", "改写条款", clause_diff["summary"]["revised_count"]),
        ("条款层", "条号顺移", clause_diff["summary"]["renumbered_count"]),
        ("参数层", "明确数值变化", parameter_diff["summary"]["changed_count"]),
        ("参数层", "新增参数项", parameter_diff["summary"]["added_count"]),
        ("参数层", "删除参数项", parameter_diff["summary"]["removed_count"]),
        ("参数层", "未变化参数项", parameter_diff["summary"]["unchanged_count"]),
        ("联动层", "条款参数联动行数", len(linkage_rows)),
    ]
    for row_idx, row in enumerate(rows, start=4):
        for col_idx, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP_ALIGNMENT

    note_row = len(rows) + 6
    write_header(sheet, note_row, ["工作簿使用说明"], fill=NOTE_FILL)
    notes = [
        "条款改写 sheet 仅保留正文主体发生变化的条款；正文主体完全一致的条目不会保留在该 sheet。",
        "正文主体一致但条号变化的条目统一进入条号顺移 sheet，因此条款改写中不应再出现相似度为 1 的误分类行。",
        "参数差异 sheet 新增关键变化、建议汇报、建议人工复核等列，便于继续做人工标注和项目复盘。",
        "遇到复杂公式或计算表达时，工作簿不直接展示乱码公式片段，而以提示语替代，详细口径需回到原文核对。",
    ]
    for idx, note in enumerate(notes, start=note_row + 1):
        cell = sheet.cell(row=idx, column=1, value=note)
        cell.alignment = WRAP_ALIGNMENT

    autosize(sheet, {1: 72, 2: 20, 3: 12})
    finalize_sheet(sheet, freeze_panes="A4")


def fill_clause_sheet(sheet, title: str, rows: List[Dict], kind: str) -> None:
    sheet.title = title
    if kind == "revised":
        header = ["旧条号", "新条号", "条款名称", "章节", "相似度", "旧文本", "新文本"]
    elif kind == "renumbered":
        header = ["旧条号", "新条号", "条款名称", "章节", "相似度"]
    else:
        header = ["条号", "条款名称", "章节", "文本"]

    write_header(sheet, 1, header)

    for row_idx, row in enumerate(rows, start=2):
        if kind == "revised":
            values = [
                row["old_marker"],
                row["new_marker"],
                row["label"] or "",
                row.get("chapter_title", ""),
                row["similarity"],
                safe_clause_text(row, "old_text"),
                safe_clause_text(row, "new_text"),
            ]
        elif kind == "renumbered":
            values = [
                row["old_marker"],
                row["new_marker"],
                row["label"] or "",
                row.get("chapter_title", ""),
                row["similarity"],
            ]
        else:
            values = [
                row["marker"],
                row["label"] or "",
                row.get("chapter_title", ""),
                safe_clause_text(row, "text"),
            ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP_ALIGNMENT

    if kind == "revised":
        autosize(sheet, {1: 10, 2: 10, 3: 22, 4: 18, 5: 10, 6: 80, 7: 80})
    elif kind == "renumbered":
        autosize(sheet, {1: 10, 2: 10, 3: 24, 4: 18, 5: 10})
    else:
        autosize(sheet, {1: 10, 2: 24, 3: 18, 4: 90})
    finalize_sheet(sheet)


def fill_parameter_sheet(sheet, parameter_rows: List[Dict]) -> None:
    sheet.title = "参数差异"
    write_header(
        sheet,
        1,
        [
            "类型",
            "变更细类",
            "参数类别",
            "适用主体",
            "是否关键变化",
            "是否建议汇报",
            "是否建议人工复核",
            "参数名称",
            "参数上下文",
            "旧值",
            "新值",
            "单位",
            "旧来源",
            "新来源",
            "业务影响",
            "证据摘录",
        ],
    )

    for row_idx, row in enumerate(parameter_rows, start=2):
        values = [
            row["change_type"],
            row["change_detail"],
            row["category_label"],
            row["subject"],
            row["is_key_change"],
            row["should_report"],
            row["needs_manual_review"],
            row["parameter_name"],
            row["context_label"],
            row["old_value"],
            row["new_value"],
            row["unit"],
            row["old_source_ref"],
            row["new_source_ref"],
            row["business_impact"],
            row["statement_text"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP_ALIGNMENT

    autosize(
        sheet,
        {
            1: 12,
            2: 12,
            3: 14,
            4: 16,
            5: 12,
            6: 12,
            7: 14,
            8: 28,
            9: 30,
            10: 18,
            11: 18,
            12: 10,
            13: 36,
            14: 36,
            15: 34,
            16: 72,
        },
    )
    finalize_sheet(sheet)


def fill_linkage_sheet(sheet, rows: List[Dict]) -> None:
    sheet.title = "条款参数联动"
    write_header(
        sheet,
        1,
        [
            "章节",
            "旧条号",
            "新条号",
            "条款名称",
            "参数名称",
            "变更类型",
            "变更细类",
            "参数类别",
            "适用主体",
            "是否关键变化",
            "是否建议人工复核",
            "参数旧值",
            "参数新值",
            "参数来源",
            "联动方式",
            "业务影响",
        ],
        fill=SUBHEADER_FILL,
    )

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row["chapter_title"],
            row["clause_marker_old"],
            row["clause_marker_new"],
            row["clause_label"],
            row["parameter_name"],
            row["change_type"],
            row["change_detail"],
            row["category_label"],
            row["subject"],
            row["is_key_change"],
            row["needs_manual_review"],
            row["parameter_old_value"],
            row["parameter_new_value"],
            row["parameter_source_ref"],
            row["link_basis"],
            row["business_impact"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP_ALIGNMENT

    autosize(
        sheet,
        {
            1: 18,
            2: 10,
            3: 10,
            4: 24,
            5: 26,
            6: 12,
            7: 12,
            8: 14,
            9: 16,
            10: 12,
            11: 14,
            12: 16,
            13: 16,
            14: 40,
            15: 16,
            16: 34,
        },
    )
    finalize_sheet(sheet)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build Excel workbook for clause diff + parameter diff")
    parser.add_argument("--old-structured", required=True)
    parser.add_argument("--new-structured", required=True)
    parser.add_argument("--clause-diff", required=True)
    parser.add_argument("--parameter-diff", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    old_structured = json.loads(resolve_path(args.old_structured).read_text(encoding="utf-8"))
    new_structured = json.loads(resolve_path(args.new_structured).read_text(encoding="utf-8"))
    clause_diff = json.loads(resolve_path(args.clause_diff).read_text(encoding="utf-8"))
    parameter_diff = json.loads(resolve_path(args.parameter_diff).read_text(encoding="utf-8"))
    output = resolve_path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)

    parameter_rows = collect_parameter_rows(parameter_diff)
    linkage_rows = collect_linkage_rows(old_structured, new_structured, clause_diff, parameter_diff)

    wb = Workbook()
    fill_summary(wb.active, clause_diff, parameter_diff, linkage_rows)
    fill_clause_sheet(wb.create_sheet(), "条款改写", clause_diff.get("revised_clauses", []), "revised")
    fill_clause_sheet(wb.create_sheet(), "新增条款", clause_diff.get("added_clauses", []), "plain")
    fill_clause_sheet(wb.create_sheet(), "删除条款", clause_diff.get("removed_clauses", []), "plain")
    fill_clause_sheet(wb.create_sheet(), "条号顺移", clause_diff.get("renumbered_clauses", []), "renumbered")
    fill_parameter_sheet(wb.create_sheet(), parameter_rows)
    fill_linkage_sheet(wb.create_sheet(), linkage_rows)

    wb.save(output)
    print(f"Workbook written: {output}")
    print(f"Parameter rows: {len(parameter_rows)}")
    print(f"Linked rows: {len(linkage_rows)}")
    refresh_frontend_dataset()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
